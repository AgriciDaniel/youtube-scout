import datetime as dt
import json
import sys
from pathlib import Path

import pytest

SKILL_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_DIR / "scripts"))

import scout  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
NOW = dt.datetime(2026, 9, 11, 12, 0, tzinfo=dt.timezone.utc).timestamp()
CATEGORIES = {"22": "People & Blogs", "26": "Howto & Style"}


@pytest.fixture
def videos():
    return json.loads((FIXTURES / "videos.json").read_text())["items"]


@pytest.fixture
def channels():
    data = json.loads((FIXTURES / "channels.json").read_text())["items"]
    return {c["id"]: c for c in data}


@pytest.fixture
def comments_payload():
    return json.loads((FIXTURES / "comments.json").read_text())


@pytest.fixture
def ytdlp_info():
    return json.loads((FIXTURES / "ytdlp_info.json").read_text())


@pytest.fixture
def captions():
    return json.loads((FIXTURES / "captions.json3").read_text())


@pytest.fixture
def rows(videos, channels):
    relevance = {v["id"]: i for i, v in enumerate(videos, start=1)}
    return scout.build_rows(videos, channels, CATEGORIES, relevance, now=NOW)


def by_id(rows):
    return {r["Video ID"]: r for r in rows}


# --- pure helpers ----------------------------------------------------------

@pytest.mark.parametrize("value,expected", [
    ("PT4M20S", 260.0), ("PT31S", 31.0), ("PT1H", 3600.0), ("P1DT2H", 93600.0),
    ("P2D", 172800.0), ("PT0S", 0.0), (None, None), ("garbage", None),
])
def test_parse_iso_duration(value, expected):
    assert scout.parse_iso_duration(value) == expected


def test_engagement_matches_owt_formula():
    assert scout.engagement_pct(847288, 73883, 150) == round((73883 + 150) / 847288 * 100, 2)
    assert scout.engagement_pct(0, 10, 10) == 0.0
    assert scout.engagement_pct(None, 10, 10) == 0.0
    assert scout.engagement_pct(1000, None, 5) == 0.5


def test_to_epoch():
    assert scout.to_epoch("2026-03-01T10:00:00Z") == 1772359200.0
    assert scout.to_epoch(None) is None
    assert scout.to_epoch("nope") is None


def test_parse_env_file(tmp_path):
    env = tmp_path / ".env"
    env.write_text("# comment\nexport YOUTUBE_API_KEY='abc123'\nOTHER=\"quoted value\"\n"  # pragma: allowlist secret
                   "BARE=plain\n\nNOEQUALS\n---\n")
    parsed = scout.parse_env_file(env)
    assert parsed == {"YOUTUBE_API_KEY": "abc123", "OTHER": "quoted value", "BARE": "plain"}  # pragma: allowlist secret
    assert scout.parse_env_file(tmp_path / "missing") == {}


def test_load_api_key_prefers_env_then_files(tmp_path, monkeypatch):
    first = tmp_path / "first.env"
    second = tmp_path / "second.env"
    second.write_text("YOUTUBE_API_KEY=fromsecond\n")  # pragma: allowlist secret
    monkeypatch.setenv("YOUTUBE_API_KEY", "fromenv")
    assert scout.load_api_key([first, second]) == "fromenv"
    monkeypatch.delenv("YOUTUBE_API_KEY")
    assert scout.load_api_key([first, second]) == "fromsecond"  # missing first is skipped
    first.write_text("YOUTUBE_API_KEY=fromfirst\n")  # pragma: allowlist secret
    assert scout.load_api_key([first, second]) == "fromfirst"
    with pytest.raises(scout.ScoutError) as exc:
        scout.load_api_key([tmp_path / "nope.env"])
    assert exc.value.code == scout.EXIT_NO_KEY
    assert "fromfirst" not in str(exc.value) and "nope.env" in str(exc.value)


def test_env_file_candidates(monkeypatch, tmp_path):
    monkeypatch.delenv("SCOUT_ENV_FILE", raising=False)
    assert scout.env_file_candidates() == list(scout.DEFAULT_ENV_FILES)
    monkeypatch.setenv("SCOUT_ENV_FILE", str(tmp_path / "keys.env"))
    assert scout.env_file_candidates()[0] == tmp_path / "keys.env"


def test_classify_error():
    quota = scout.classify_error(403, {"error": {"errors": [{"reason": "quotaExceeded"}]}})
    assert quota.code == scout.EXIT_QUOTA
    bad = scout.classify_error(400, {"error": {"message": "API key not valid"}})
    assert bad.code == scout.EXIT_BAD_KEY
    other = scout.classify_error(500, None)
    assert other.code == scout.EXIT_API
    disabled = scout.classify_error(403, {"error": {"errors": [{"reason": "commentsDisabled"}]}})
    assert disabled.code == scout.EXIT_API and "commentsdisabled" in disabled.reason


def test_published_after():
    now = dt.datetime(2026, 9, 11, 15, 30, tzinfo=dt.timezone.utc)
    assert scout.published_after("any", now) is None
    assert scout.published_after("today", now) == "2026-09-11T00:00:00Z"
    assert scout.published_after("week", now) == "2026-09-04T15:30:00Z"
    assert scout.published_after("year", now) == "2025-09-11T15:30:00Z"


def test_small_helpers():
    assert scout.topic_names(["https://en.wikipedia.org/wiki/Lifestyle_(sociology)",
                              "https://en.wikipedia.org/wiki/Food", "https://en.wikipedia.org/wiki/Food"]) \
        == "Lifestyle (sociology), Food"
    assert scout.yes_no("true") == "yes" and scout.yes_no(False) == "no" and scout.yes_no(None) == ""
    assert scout.video_format(31, False) == "Short"
    assert scout.video_format(31, False, vertical="no") == "Video"
    assert scout.video_format(181, False) == "Video"
    assert scout.video_format(31, True) == "Live"
    assert scout.age_days(NOW - 86400 * 10, NOW) == 10.0
    assert scout.age_days(NOW, NOW) == 0.25
    assert scout.region_summary({"regionRestriction": {"blocked": ["DE", "FR"]}}) == "2 blocked"
    assert scout.region_summary({}) == ""
    assert scout.location_summary({"location": {"latitude": 1.23456, "longitude": 2.0}}) == "1.2346,2.0000"


# --- row building ----------------------------------------------------------

def test_build_rows_shape(rows):
    assert len(rows) == 3
    assert list(rows[0].keys()) == scout.ALL_COLUMNS
    a = by_id(rows)["vid00000001"]
    assert a["Platform"] == "YouTube"
    assert a["Creator"] == "Tea Lab" and a["Handle"] == "tealab"
    assert a["Channel URL"] == "https://www.youtube.com/@tealab"
    assert a["Video Link"] == "https://www.youtube.com/watch?v=vid00000001"
    assert a["Posted"] == 1772359200.0 and a["Published"] == "2026-03-01"
    assert a["Views"] == 120000 and a["Likes"] == 8000 and a["Comments"] == 400
    assert a["Shares"] is None
    assert a["Engagement %"] == 7.0 and a["Like %"] == 6.67 and a["Comment %"] == 0.33
    assert a["Duration s"] == 260.0 and a["Format"] == "Video"
    assert a["Relevance #"] == 1
    assert a["Age d"] == 194.08 and a["Views/day"] == 618
    assert a["Subscribers"] == 45000 and a["Views/Sub"] == 2.67
    assert a["Channel Videos"] == 150 and a["Channel Views"] == 3000000 and a["Avg Views/Video"] == 20000
    assert a["Channel Country"] == "JP" and a["Channel Created"] == "2019-06-01"
    assert a["Channel Keywords"].startswith("matcha tea") and a["Channel Topics"] == "Food"
    assert a["Category"] == "Howto & Style"
    assert a["Tags"] == "matcha, Matcha Latte, tea" and a["Tag Count"] == 3
    assert a["Language"] == "en" and a["Captions"] == "yes" and a["HD"] == "yes"
    assert a["Licensed"] == "yes" and a["Embeddable"] == "yes" and a["Made for Kids"] == "no"
    assert a["Paid Promotion"] == "yes" and a["AI Disclosure"] == "yes"
    assert a["License"] == "youtube" and a["Live"] == ""
    assert a["Topics"] == "Food, Lifestyle (sociology)"
    assert a["Blocked Regions"] == ""
    assert a["Description"] == ("Full recipe at https://tealab.example/recipe and my whisk "
                                "https://shop.example/whisk Enjoy!")
    assert a["Desc Links"] == 2 and a["Location"] == "Kyoto, Japan"
    assert a["Thumbnail"].endswith("hqdefault.jpg")
    for manual in ("Spark Code", "Local File", "Sheet Views", "Ad Status", "Top Comment",
                   "Hook", "Vertical", "Thumbnail File"):
        assert a[manual] is None


def test_edge_rows(rows):
    r = by_id(rows)
    b = r["vid00000002"]
    assert b["Subscribers"] is None and b["Views/Sub"] is None  # hidden
    assert b["Thumbnail"].endswith("maxresdefault.jpg")
    assert b["Format"] == "Short" and b["HD"] == "no" and b["Captions"] == "no"
    assert b["Blocked Regions"] == "2 blocked" and b["License"] == "creativeCommon"
    assert b["Embeddable"] == "no" and b["Paid Promotion"] == "no" and b["AI Disclosure"] == ""
    assert b["Category"] == "People & Blogs" and b["Tag Count"] == 0 and b["Desc Links"] == 0
    assert b["Avg Views/Video"] is None
    c = r["vid00000003"]
    assert c["Handle"] == "" and c["Channel URL"] == "https://www.youtube.com/channel/chanC"
    assert c["Comments"] is None and c["Engagement %"] == 18.0
    assert c["Duration s"] == 93600.0 and c["Format"] == "Live" and c["Live"] == "was live"
    assert c["Category"] == "99" and c["Blocked Regions"] == "only 1 allowed"
    assert c["HD"] == "" and c["Location"] == ""


def test_sort_rows(rows):
    ids = lambda key: [r["Video ID"] for r in scout.sort_rows(rows, key)]  # noqa: E731
    assert ids("views") == ["vid00000002", "vid00000001", "vid00000003"]
    assert ids("engagement") == ["vid00000003", "vid00000002", "vid00000001"]
    assert ids("likes") == ["vid00000002", "vid00000001", "vid00000003"]
    assert ids("recent") == ["vid00000002", "vid00000001", "vid00000003"]
    assert ids("momentum") == ["vid00000002", "vid00000001", "vid00000003"]
    assert ids("breakout") == ["vid00000001", "vid00000002", "vid00000003"]


# --- comments ---------------------------------------------------------------

def test_enrich_comments(monkeypatch, rows, comments_payload):
    calls = []

    def fake_get(endpoint, params, api_key):
        calls.append(params["videoId"])
        if params["videoId"] == "vid00000002":
            raise scout.classify_error(403, {"error": {"errors": [{"reason": "commentsDisabled"}]}})
        if params["videoId"] == "vid00000003":
            return {"items": []}
        return comments_payload

    monkeypatch.setattr(scout, "yt_get", fake_get)
    quota = scout.Quota()
    comment_rows, warnings = scout.enrich_comments(rows, 20, "k", quota)
    r = by_id(rows)
    assert warnings == [] and quota.units == 2  # disabled call costs nothing here
    assert r["vid00000001"]["Top Comment"] == "This whisk changed my mornings"
    assert r["vid00000001"]["Top Comment Likes"] == 250 and r["vid00000001"]["Comments Off"] == "no"
    assert r["vid00000002"]["Comments Off"] == "yes" and r["vid00000002"]["Top Comment"] is None
    assert r["vid00000003"]["Comments Off"] == "no"
    assert len(comment_rows) == 2
    assert list(comment_rows[0].keys()) == scout.COMMENT_COLUMNS
    assert comment_rows[0]["Author"] == "Ana" and comment_rows[0]["Replies"] == 4


def test_quota_summary():
    q = scout.Quota()
    assert q.summary() == "0 search calls of 100 a day, 0 units of 10,000"
    q.searches, q.units = 1, 3
    assert q.summary() == "1 search call of 100 a day, 3 units of 10,000"


def test_search_ids_counts_search_calls(monkeypatch):
    pages = [{"items": [{"id": {"videoId": f"v{i}"}} for i in range(50)], "nextPageToken": "p2"},
             {"items": [{"id": {"videoId": f"w{i}"}} for i in range(30)]}]
    monkeypatch.setattr(scout, "yt_get", lambda ep, params, key: pages.pop(0))
    q = scout.Quota()
    ids = scout.search_ids("matcha", 80, "any", "any", "k", q)
    assert len(ids) == 80 and q.searches == 2 and q.units == 0


def test_fetch_comments_propagates_quota(monkeypatch):
    def fake_get(endpoint, params, api_key):
        raise scout.classify_error(403, {"error": {"errors": [{"reason": "quotaExceeded"}]}})
    monkeypatch.setattr(scout, "yt_get", fake_get)
    with pytest.raises(scout.ScoutError) as exc:
        scout.fetch_comments("x", 5, "k", scout.Quota())
    assert exc.value.code == scout.EXIT_QUOTA


# --- hooks (yt-dlp) ---------------------------------------------------------

def test_pick_caption_url(ytdlp_info):
    lang, url = scout.pick_caption_url(ytdlp_info)
    assert (lang, url) == ("en", "https://captions.example/vid2.json3")
    manual = {"subtitles": {"en-GB": [{"ext": "json3", "url": "m"}]}, "automatic_captions": ytdlp_info["automatic_captions"]}
    assert scout.pick_caption_url(manual) == ("en-GB", "m")
    only_orig = {"automatic_captions": {"ja-orig": [{"ext": "json3", "url": "o"}]}}
    assert scout.pick_caption_url(only_orig) == ("ja-orig", "o")
    assert scout.pick_caption_url({}) == ("", "")


def test_parse_json3(captions):
    hook, full = scout.parse_json3(captions, 15)
    assert hook == "make it or buy it iced matcha latte at home"
    assert full == "make it or buy it iced matcha latte at home costs six dollars worth it"


def test_heatmap_hotspots(ytdlp_info):
    # 12s is a true peak (0.9 above 0.2 and 0.85); the opening decay and the 25s tail are not.
    assert scout.heatmap_hotspots(ytdlp_info["heatmap"]) == [12]
    assert scout.heatmap_hotspots(None) == [] and scout.heatmap_hotspots([]) == []
    decay_only = [{"start_time": i, "end_time": i + 1, "value": 1 - i / 10} for i in range(10)]
    assert scout.heatmap_hotspots(decay_only) == []
    two_peaks = decay_only + [{"start_time": 20, "end_time": 21, "value": 0.7},
                              {"start_time": 21, "end_time": 22, "value": 0.1},
                              {"start_time": 40, "end_time": 41, "value": 0.9},
                              {"start_time": 41, "end_time": 42, "value": 0.2}]
    assert scout.heatmap_hotspots(two_peaks) == [40, 20]


def test_probe_video_and_enrich_hooks(rows, ytdlp_info, captions, tmp_path):
    fetched = []
    runs = []

    def runner(url):
        runs.append(url)
        return ytdlp_info if "vid00000002" in url else None

    def fetcher(url):
        fetched.append(url)
        return captions

    cache = tmp_path / "cache"
    url2 = "https://www.youtube.com/watch?v=vid00000002"
    info = scout.probe_video(url2, 15, runner, fetcher, cache_dir=cache)
    assert info["Vertical"] == "yes" and info["FPS"] == 60 and info["Chapters"] == 0
    assert info["Most Replayed s"] == 12 and info["Replay Hotspots"] == "12s"
    assert info["Hook"].startswith("make it or buy it") and info["Transcript Words"] == 15
    assert info["Language"] == "en"
    assert fetched == ["https://captions.example/vid2.json3"]
    assert (cache / "vid00000002.json").exists()
    assert scout.probe_video("x", 15, lambda u: None, fetcher, cache_dir=cache) == {}

    # Second probe is served entirely from cache: no yt-dlp run, no caption fetch.
    again = scout.probe_video(url2, 15, runner, fetcher, cache_dir=cache)
    assert again["Hook"] == info["Hook"] and len(runs) == 1 and len(fetched) == 1

    sleeps = []
    transcript_rows, warnings = scout.enrich_hooks(
        rows, 15, prober=lambda url, s, fetch=True: scout.probe_video(url, s, runner, fetcher, cache_dir=cache),
        delay_s=1.5, sleeper=sleeps.append)
    r = by_id(rows)
    assert r["vid00000002"]["Hook"].startswith("make it") and r["vid00000002"]["Vertical"] == "yes"
    assert r["vid00000002"]["Language"] == "en"
    assert len(warnings) == 2 and all("probe failed" in w for w in warnings)
    assert sleeps == [1.5, 1.5]  # paced between videos, not before the first
    assert len(transcript_rows) == 1
    assert list(transcript_rows[0].keys()) == scout.TRANSCRIPT_COLUMNS
    assert transcript_rows[0]["Transcript"].endswith("worth it")


def test_probe_video_rate_limit_backoff(ytdlp_info, captions, tmp_path):
    attempts = []
    sleeps = []

    def flaky(url):
        attempts.append(url)
        if len(attempts) < 3:
            raise scout.RateLimited()
        return captions

    url2 = "https://www.youtube.com/watch?v=vid00000002"
    info = scout.probe_video(url2, 15, lambda u: ytdlp_info, flaky, cache_dir=tmp_path,
                             sleeper=sleeps.append)
    assert info["Hook"].startswith("make it") and "_rate_limited" not in info
    assert sleeps == list(scout.CAPTION_BACKOFF_S) and len(attempts) == 3

    def always(url):
        raise scout.RateLimited()

    info = scout.probe_video("https://www.youtube.com/watch?v=vid00000099", 15,
                             lambda u: ytdlp_info, always, cache_dir=tmp_path, sleeper=lambda s: None)
    assert info["_rate_limited"] is True and "Hook" not in info and info["Vertical"] == "yes"
    # metadata cached, captions not, so a later run retries only the captions
    cached = scout.cache_read("vid00000099", tmp_path)
    assert cached["info"]["caption_url"] and "captions" not in cached

    rows = [{"Video ID": v, "Video Link": f"https://www.youtube.com/watch?v={v}",
             "Handle": "h", "Title": "t", "Format": "Short", "Language": ""}
            for v in ("vid00000099", "vid00000098", "vid00000097")]
    fetch_flags = []

    def prober(url, s, fetch=True):
        fetch_flags.append(fetch)
        return scout.probe_video(url, s, lambda u: ytdlp_info, always, cache_dir=tmp_path,
                                 sleeper=lambda x: None, fetch_captions=fetch)

    messages = []
    _, warnings = scout.enrich_hooks(rows, 15, progress=messages.append, prober=prober, delay_s=0)
    assert fetch_flags == [True, False, False]  # breaker trips after the first full failure
    assert len(warnings) == 1 and "rate-limited by YouTube for 3 videos" in warnings[0]
    assert any("skipping caption fetches" in m for m in messages)
    assert all(scout.cache_read(r["Video ID"], tmp_path)["info"]["fps"] == 60 for r in rows)


def test_hooks_demotes_horizontal_short(rows):
    def prober(url, s, fetch=True):
        return {"Vertical": "no"} if "vid00000002" in url else {}
    scout.enrich_hooks(rows, 15, prober=prober, delay_s=0)
    assert by_id(rows)["vid00000002"]["Format"] == "Video"


# --- channels and summary ---------------------------------------------------

def test_build_channel_rows(rows):
    ch = scout.build_channel_rows(rows)
    assert [c["Handle"] for c in ch] == ["quicksips", "tealab", ""]
    assert list(ch[0].keys()) == scout.CHANNEL_COLUMNS
    tea = next(c for c in ch if c["Handle"] == "tealab")
    assert tea["Videos in Sample"] == 1 and tea["Sample Views"] == 120000
    assert tea["Best Video"] == "How to whisk matcha properly" and tea["Best Video Views"] == 120000
    assert tea["Sample Avg Eng %"] == 7.0 and tea["Channel Country"] == "JP"


def test_build_summary(rows):
    args = scout.build_parser().parse_args(["matcha", "--comments", "--hooks"])
    pairs = dict(scout.build_summary("matcha", args, scout.sort_rows(rows, "views"), "1 search call of 100 a day, 5 units of 10,000",
                                     now=dt.datetime(2026, 9, 11, 12, 0)))
    assert pairs["Topic"] == "matcha" and pairs["Videos"] == 3 and pairs["Quota used"] == "1 search call of 100 a day, 5 units of 10,000"
    assert pairs["Filters"] == "max 50, since any, length any, sort views, comments 20, hooks 15.0s"
    assert pairs["Total views"] == 1025000 and pairs["Median views"] == 120000
    assert pairs["Shorts share"] == "33%" and pairs["Paid promotion share"] == "33%"
    assert pairs["Published range"] == "2025-11-20 to 2026-05-15"
    assert pairs["Top video"].startswith("Iced matcha latte") and "900,000" in pairs["Top video"]
    assert pairs["Top tags"] == "matcha (1), matcha latte (1), tea (1)"
    assert pairs["Categories"] == "People & Blogs (1), Howto & Style (1), 99 (1)"


# --- workbook round trip ---------------------------------------------------

def test_write_then_append(tmp_path, rows):
    from openpyxl import load_workbook

    out = tmp_path / "scout.xlsx"
    sorted_rows = scout.sort_rows(rows, "views")
    scout.write_workbook(sorted_rows, out, channel_rows=scout.build_channel_rows(sorted_rows),
                         comment_rows=[{"Video ID": "vid00000001", "Comment": "hi"}],
                         transcript_rows=[{"Video ID": "vid00000002", "Transcript": "t"}],
                         summary=[("Topic", "x"), ("Videos", 3)])
    wb = load_workbook(out)
    assert wb.sheetnames == [scout.SHEET_NAME, scout.CHANNELS_SHEET, scout.COMMENTS_SHEET,
                             scout.TRANSCRIPTS_SHEET, scout.SUMMARY_SHEET]
    ws = wb[scout.SHEET_NAME]
    header = [c.value for c in ws[1]]
    assert header == [c for c in scout.ALL_COLUMNS if c not in scout.OPTIONAL_COLUMNS]
    assert "Shares" not in header and "Spark Code" not in header and "Hook" not in header
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    assert ws["A1"].fill.fgColor.rgb == scout.HEADER_FILL
    assert ws["A1"].font.bold and ws["A1"].font.name == "Arial"
    assert ws["G2"].number_format == "#,##0" and ws["J2"].number_format == '0.00"%"'
    assert ws["D2"].value == "vid00000002" and ws.max_row == 4
    assert ws["E2"].hyperlink.target == "https://www.youtube.com/watch?v=vid00000002"
    col = header.index("Channel URL") + 1
    assert ws.cell(2, col).hyperlink is not None
    assert wb[scout.SUMMARY_SHEET]["A2"].value == "Videos" and wb[scout.SUMMARY_SHEET]["B2"].value == 3
    assert wb[scout.CHANNELS_SHEET]["A1"].value == "Handle"

    new = dict(rows[0])
    new.update({"Video ID": "vid00000009", "Video Link": "https://www.youtube.com/watch?v=vid00000009",
                "Views": 300000, "Likes": 1, "Comments": 1, "Engagement %": 0.0})
    path, added, skipped = scout.append_workbook(
        [rows[0], rows[1], new], out, "views",
        comment_rows=[{"Video ID": "vid00000001", "Comment": "dup"}, {"Video ID": "vid00000009", "Comment": "new"}],
        transcript_rows=[{"Video ID": "vid00000009", "Transcript": "t9"}])
    assert (added, skipped) == (1, 2)
    assert len(list(tmp_path.glob("scout.backup-*.xlsx"))) == 1
    wb = load_workbook(out)
    ws = wb[scout.SHEET_NAME]
    ids = [ws.cell(r, 4).value for r in range(2, ws.max_row + 1)]
    assert ids == ["vid00000002", "vid00000009", "vid00000001", "vid00000003"]
    assert ws.freeze_panes == "A2"
    comments = wb[scout.COMMENTS_SHEET]
    assert [comments.cell(r, 1).value for r in range(2, comments.max_row + 1)] == ["vid00000001", "vid00000009"]
    transcripts = wb[scout.TRANSCRIPTS_SHEET]
    assert transcripts.max_row == 3

    _, added, skipped = scout.append_workbook([rows[0], rows[1], new], out, "views")
    assert (added, skipped) == (0, 3)


def test_append_into_owt_style_sheet(tmp_path, rows):
    """Existing sheet with only the 16 OWT columns, string-typed numbers, and a hyperlink."""
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "owt.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = scout.SHEET_NAME
    ws.append(scout.OWT_COLUMNS)
    ws.append(["TikTok", "SOMEONE", "someone", "777", "https://t/777", "1775588382.0",
               847288.0, 73883.0, 150.0, 1214.0, 8.88, 45.2, "", "someone_777.mp4", 298800.0, ""])
    ws["E2"].hyperlink = "https://t/777?share=1"
    ws.append(["Instagram", "OTHER", "", "", "https://i/x", "46109.0", 5070.0, "", "", "",
               0.0, "", "5070", "", 5070.0, ""])
    ws.column_dimensions["A"].width = 9.9
    wb.save(path)

    _, added, skipped = scout.append_workbook(rows, path, "views",
                                              comment_rows=[{"Video ID": "vid00000001", "Comment": "c"}])
    assert (added, skipped) == (3, 0)
    wb = load_workbook(path)
    ws = wb[scout.SHEET_NAME]
    assert [c.value for c in ws[1]] == scout.OWT_COLUMNS + scout.EXTRA_COLUMNS
    assert abs(ws.column_dimensions["A"].width - 9.9) < 1e-6
    body = [[ws.cell(r, c).value for c in range(1, 17)] for r in range(2, ws.max_row + 1)]
    assert [b[6] for b in body] == [900000, 847288.0, 120000, 5070.0, 5000]
    tiktok = next(b for b in body if b[0] == "TikTok")
    assert tiktok[13] == "someone_777.mp4" and tiktok[14] == 298800.0
    assert ws.max_row == 6
    links = {ws.cell(r, 5).value: (ws.cell(r, 5).hyperlink.target if ws.cell(r, 5).hyperlink else None)
             for r in range(2, ws.max_row + 1)}
    assert links["https://t/777"] == "https://t/777?share=1"
    assert links["https://www.youtube.com/watch?v=vid00000002"] == "https://www.youtube.com/watch?v=vid00000002"
    assert links["https://i/x"] is None
    assert wb[scout.COMMENTS_SHEET]["A2"].value == "vid00000001"


def test_new_workbook_columns(rows):
    cols = scout.new_workbook_columns(rows)
    assert cols[:9] == ["Platform", "Creator", "Handle", "Video ID", "Video Link", "Posted",
                        "Views", "Likes", "Comments"]
    assert "Engagement %" in cols and "Local File" not in cols
    rows[0]["Local File"] = "x.mp4"
    rows[1]["Hook"] = "hi"
    cols = scout.new_workbook_columns(rows)
    assert "Local File" in cols and "Hook" in cols and "Spark Code" not in cols


def test_save_errors(tmp_path, rows, monkeypatch):
    from openpyxl import Workbook
    wb = Workbook()
    target = tmp_path / "locked.xlsx"

    def denied(path):
        raise PermissionError(13, "Permission denied", str(path))
    monkeypatch.setattr(wb, "save", denied)
    with pytest.raises(scout.ScoutError) as exc:
        scout._save(wb, target)
    assert exc.value.code == scout.EXIT_FILE and "open in another program" in str(exc.value)


def test_yt_get_retries_once_on_server_error(monkeypatch):
    import io
    import urllib.error
    calls = []
    sleeps = []

    def fake_urlopen(req, timeout):
        calls.append(req.full_url)
        if len(calls) == 1:
            raise urllib.error.HTTPError(req.full_url, 503, "unavailable", {}, io.BytesIO(b"{}"))
        return io.BytesIO(b'{"items": []}')

    monkeypatch.setattr(scout.urllib.request, "urlopen", fake_urlopen)
    assert scout.yt_get("videos", {"id": "x"}, "k", sleeper=sleeps.append) == {"items": []}
    assert len(calls) == 2 and sleeps == [2.0]
    assert "key=k" in calls[0]

    calls.clear()

    def always_down(req, timeout):
        calls.append(1)
        raise urllib.error.URLError("dns")
    monkeypatch.setattr(scout.urllib.request, "urlopen", always_down)
    with pytest.raises(scout.ScoutError) as exc:
        scout.yt_get("videos", {"id": "x"}, "k", sleeper=lambda s: None)
    assert exc.value.code == scout.EXIT_API and len(calls) == 2


def test_append_missing_file(tmp_path, rows):
    with pytest.raises(scout.ScoutError) as exc:
        scout.append_workbook(rows, tmp_path / "nope.xlsx", "views")
    assert exc.value.code == scout.EXIT_FILE


# --- CLI ---------------------------------------------------------------------

def test_cli_usage_errors():
    assert scout.run(["matcha", "--max", "0"]) == scout.EXIT_USAGE
    assert scout.run(["matcha", "--out", "a.xlsx", "--into", "b.xlsx"]) == scout.EXIT_USAGE
    assert scout.run(["matcha", "--comments", "101"]) == scout.EXIT_USAGE


def test_cli_end_to_end_offline(monkeypatch, tmp_path, videos, channels, comments_payload,
                                ytdlp_info, captions, capsys):
    from openpyxl import load_workbook

    monkeypatch.setenv("YOUTUBE_API_KEY", "test")
    monkeypatch.chdir(tmp_path)
    calls = []

    def fake_get(endpoint, params, api_key):
        calls.append(endpoint)
        if endpoint == "search":
            return {"items": [{"id": {"videoId": v["id"]}} for v in videos]}
        if endpoint == "videos":
            return {"items": videos}
        if endpoint == "channels":
            return {"items": list(channels.values())}
        if endpoint == "videoCategories":
            return {"items": [{"id": k, "snippet": {"title": v}} for k, v in CATEGORIES.items()]}
        if endpoint == "commentThreads":
            return comments_payload if params["videoId"] == "vid00000001" else {"items": []}
        raise AssertionError(endpoint)

    monkeypatch.setattr(scout, "yt_get", fake_get)
    monkeypatch.setattr(scout, "_run_ytdlp_json", lambda url: ytdlp_info if "vid00000002" in url else None)
    monkeypatch.setattr(scout, "_fetch_json", lambda url: captions)
    monkeypatch.setattr(scout, "CACHE_DIR", tmp_path / "cache")
    monkeypatch.setattr(scout, "CAPTION_DELAY_S", 0.0)
    monkeypatch.setattr(scout.time, "sleep", lambda s: None)

    assert scout.run(["matcha", "recipe", "--dry-run"]) == scout.EXIT_OK
    out, err = capsys.readouterr()
    assert calls == ["search", "videos", "channels", "videoCategories"]
    assert "quicksips" in out and "900,000" in out and "Short" in out
    assert "Quota used: 1 search call of 100 a day, 3 units of 10,000." in err
    assert not list(tmp_path.glob("*.xlsx"))

    calls.clear()
    assert scout.run(["matcha", "recipe", "--comments", "5", "--hooks"]) == scout.EXIT_OK
    out, err = capsys.readouterr()
    assert calls.count("commentThreads") == 3
    assert "Quota used: 1 search call of 100 a day, 6 units of 10,000." in err
    assert "warning: vid00000001: yt-dlp probe failed" in err
    written = list(tmp_path.glob("scout-matcha-recipe-*.xlsx"))
    assert len(written) == 1
    wb = load_workbook(written[0])
    assert wb.sheetnames == ["Videos", "Channels", "Comments", "Transcripts", "Summary"]
    ws = wb["Videos"]
    header = [c.value for c in ws[1]]
    assert "Hook" in header and "Top Comment" in header and "Shares" not in header
    row2 = dict(zip(header, [c.value for c in ws[2]], strict=False))
    assert row2["Video ID"] == "vid00000002" and row2["Hook"].startswith("make it")
    assert row2["Vertical"] == "yes" and row2["Most Replayed s"] == 12
    row3 = dict(zip(header, [c.value for c in ws[3]], strict=False))
    assert row3["Top Comment"] == "This whisk changed my mornings" and row3["Top Comment Likes"] == 250
    assert wb["Comments"].max_row == 3 and wb["Transcripts"].max_row == 2


def test_default_out_path():
    p = scout.default_out_path("Matcha Recipe!! 2026", dt.date(2026, 9, 11))
    assert p.name == "scout-matcha-recipe-2026-2026-09-11.xlsx"


# --- local transcription (--transcribe) ---------------------------------------

FAKE_RESULT = {
    "backend": "whisper turbo (gpu)",
    "language": "en",
    "segments": [[0.0, 5.7, "Best SEO tools for 2025."], [5.7, 20.0, "Number one is free and it ranks."]],
    "words": [[0.0, " Best"], [0.4, " SEO"], [0.8, " tools"], [1.2, " for"], [1.6, " 2025."],
              [5.7, " Number"], [14.9, " one"], [15.2, " is"], [16.0, " free"]],
}


def test_hook_from_transcript():
    assert scout.hook_from_transcript(FAKE_RESULT, 15) == "Best SEO tools for 2025. Number one"
    assert scout.hook_from_transcript(FAKE_RESULT, 1) == "Best SEO tools"
    no_words = {"segments": FAKE_RESULT["segments"], "words": []}
    assert scout.hook_from_transcript(no_words, 5) == "Best SEO tools for 2025."


def test_enrich_transcribe_fills_missing_and_labels_source(tmp_path, rows):
    media = tmp_path / "downloads"
    media.mkdir()
    (media / "tealab_vid00000001.mp4").write_bytes(b"x")
    rows = scout.sort_rows(rows, "views")
    r = by_id(rows)
    r["vid00000001"]["Local File"] = "tealab_vid00000001.mp4"
    r["vid00000002"].update({"Hook": "from captions", "Transcript Words": 12, "Transcript Source": "captions"})
    calls, fetched = [], []

    def transcriber(path):
        calls.append(path.name)
        return FAKE_RESULT

    def fetcher(row):
        fetched.append(row["Video ID"])
        f = tmp_path / f"{row['Video ID']}.m4a"
        f.write_bytes(b"a")
        return f

    cache = tmp_path / "cache"
    trows, warnings = scout.enrich_transcribe(rows, 15, transcriber, media_dir=media,
                                              cache_dir=cache, audio_fetcher=fetcher)
    assert warnings == []
    assert calls == ["tealab_vid00000001.mp4", "vid00000003.m4a"]  # local mp4 used first
    assert fetched == ["vid00000003"]                               # audio fetched only when needed
    a, b = r["vid00000001"], r["vid00000002"]
    assert a["Hook"] == "Best SEO tools for 2025. Number one" and a["Transcript Words"] == 12
    assert a["Transcript Source"] == "whisper" and a["Language"] == "en"
    assert b["Hook"] == "from captions" and b["Transcript Source"] == "captions"  # untouched
    assert [t["Video ID"] for t in trows] == ["vid00000001", "vid00000003"]
    assert list(trows[0].keys()) == scout.TRANSCRIPT_COLUMNS
    assert trows[0]["Source"] == "whisper (whisper turbo (gpu))"

    # second pass reuses the cache: no transcription, no audio fetch
    for key in ("vid00000001", "vid00000003"):
        r[key]["Transcript Words"] = None
    calls.clear()
    fetched.clear()
    scout.enrich_transcribe(rows, 15, transcriber, media_dir=media, cache_dir=cache, audio_fetcher=fetcher)
    assert calls == [] and fetched == []


def test_enrich_transcribe_warnings(tmp_path, rows):
    def boom(path):
        raise RuntimeError("cuda")
    trows, warnings = scout.enrich_transcribe(rows, 15, boom, cache_dir=tmp_path,
                                              audio_fetcher=lambda row: None)
    assert trows == [] and len(warnings) == 3 and all("no audio" in w for w in warnings)
    f = tmp_path / "a.m4a"
    f.write_bytes(b"a")
    trows, warnings = scout.enrich_transcribe(rows, 15, boom, cache_dir=tmp_path,
                                              audio_fetcher=lambda row: f)
    assert trows == [] and all("transcription failed (RuntimeError)" in w for w in warnings)
    silent = {"backend": "x", "language": "", "segments": [], "words": []}
    scout.enrich_transcribe(rows, 15, lambda p: silent, cache_dir=tmp_path / "c2", audio_fetcher=lambda row: f)
    assert all(r["Transcript Source"] == "no speech" and not r["Hook"] for r in rows)


def test_fetch_audio_reuses_existing_and_handles_failure(tmp_path):
    (tmp_path / "abc123.m4a").write_bytes(b"a")
    assert scout.fetch_audio("u", "abc123", tmp_path, runner=lambda c: 1 / 0) == tmp_path / "abc123.m4a"

    def ok(cmd):
        assert "bestaudio[ext=m4a]/bestaudio" in cmd
        (tmp_path / "zzz999.webm.part").write_bytes(b"partial")
        (tmp_path / "zzz999.webm").write_bytes(b"a")
    assert scout.fetch_audio("u", "zzz999", tmp_path, runner=ok) == tmp_path / "zzz999.webm"

    def fail(cmd):
        raise scout.subprocess.CalledProcessError(1, cmd)
    assert scout.fetch_audio("u", "nope00", tmp_path, runner=fail) is None


def test_transcriber_without_backends(monkeypatch):
    import builtins
    real_import = builtins.__import__

    def fake_import(name, *a, **k):
        if name in ("torch", "whisper", "faster_whisper"):
            raise ImportError(name)
        return real_import(name, *a, **k)
    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(scout.ScoutError) as exc:
        scout.Transcriber()(scout.Path("x.m4a"))
    assert exc.value.code == scout.EXIT_FILE and "--transcribe" in str(exc.value)


def test_download_limit(tmp_path, monkeypatch):
    rows = [{"Video ID": f"v{i}", "Handle": "h", "Thumbnail": f"https://t/{i}.jpg",
             "Video Link": f"https://www.youtube.com/watch?v=v{i}"} for i in range(4)]
    thumbs, videos = [], []
    monkeypatch.setattr(scout, "_download_file", lambda url, path: thumbs.append(path.name) or path.write_bytes(b"j") or True)

    def fake_run(cmd, **kw):
        videos.append(cmd[-1])
        out = cmd[cmd.index("-o") + 1].replace("%(ext)s", "mp4")
        scout.Path(out).write_bytes(b"v")
    monkeypatch.setattr(scout.subprocess, "run", fake_run)
    assert scout.download_videos(rows, tmp_path, 2) == []
    assert thumbs == [f"h_v{i}.jpg" for i in range(4)]            # every thumbnail
    assert sorted(v[-2:] for v in videos) == ["v0", "v1"]       # only the top 2 videos
    assert [r.get("Local File") for r in rows] == ["h_v0.mp4", "h_v1.mp4", None, None]
    assert all(r["Thumbnail File"] for r in rows)
    parser = scout.build_parser()
    assert parser.parse_args(["x"]).download is None
    assert parser.parse_args(["x", "--download"]).download == 0
    assert parser.parse_args(["x", "--download", "10"]).download == 10


def test_download_video_uses_merged_format(tmp_path, monkeypatch):
    seen = {}

    def fake_run(cmd, **kw):
        seen["cmd"], seen["timeout"] = cmd, kw["timeout"]
        scout.Path(cmd[cmd.index("-o") + 1].replace("%(ext)s", "mp4")).write_bytes(b"v")
    monkeypatch.setattr(scout.subprocess, "run", fake_run)
    row = {"Video ID": "abc", "Handle": "h", "Video Link": "https://www.youtube.com/watch?v=abc",
           "Duration s": 5216.0}
    assert scout._download_video(row, tmp_path) == ("h_abc.mp4", None)
    assert seen["cmd"][seen["cmd"].index("-f") + 1] == scout.DOWNLOAD_FORMAT
    assert "--merge-output-format" in seen["cmd"] and "height<=480" in scout.DOWNLOAD_FORMAT
    assert seen["timeout"] == 5216.0 * 1.5          # long video, long timeout
    assert scout._media_timeout({"Duration s": 30}, 600) == 600


def test_enrich_transcribe_retries_with_audio_when_local_file_fails(tmp_path, rows):
    media = tmp_path / "downloads"
    media.mkdir()
    (media / "silent.mp4").write_bytes(b"x")
    row = rows[0]
    row["Local File"] = "silent.mp4"
    audio = tmp_path / "a.m4a"
    audio.write_bytes(b"a")

    def transcriber(path):
        if path.name == "silent.mp4":
            raise RuntimeError("Output file does not contain any stream")
        return FAKE_RESULT
    trows, warnings = scout.enrich_transcribe([row], 15, transcriber, media_dir=media,
                                              cache_dir=tmp_path / "c", audio_fetcher=lambda r: audio)
    assert warnings == [] and row["Transcript Source"] == "whisper" and len(trows) == 1


def test_hooks_warning_mentions_transcribe_when_active(rows):
    def prober(url, s, fetch=True):
        return {"Vertical": "yes", "_rate_limited": True}
    _, w = scout.enrich_hooks(rows, 15, prober=prober, delay_s=0, transcribe_after=True)
    assert len(w) == 1 and "--transcribe fills them locally" in w[0]
    _, w = scout.enrich_hooks(rows, 15, prober=prober, delay_s=0)
    assert "Add --transcribe" in w[0]


def test_fit_whisper_model():
    assert scout.fit_whisper_model("turbo", 12.0) == "turbo"
    assert scout.fit_whisper_model("turbo", 2.4) == "small"      # busy GPU steps down
    assert scout.fit_whisper_model("turbo", 1.3) == "tiny"
    assert scout.fit_whisper_model("turbo", 0.5) is None
    assert scout.fit_whisper_model("small", 2.4) == "small"


def test_transcriber_falls_back_to_cpu_on_oom(monkeypatch):
    t = scout.Transcriber()

    class GpuModel:
        def transcribe(self, *a, **k):
            raise RuntimeError("CUDA out of memory. Tried to allocate 20.00 MiB")

    class CpuModel:
        def transcribe(self, *a, **k):
            return {"language": "en", "segments": [{"start": 0.0, "end": 2.0, "text": " hi there",
                                                    "words": [{"start": 0.0, "word": " hi"}]}]}

    t._m, t._kind, t.backend = GpuModel(), "openai", "whisper turbo (gpu)"

    def fake_cpu():
        t._m, t._kind, t.backend = CpuModel(), "openai-cpu", "whisper small (cpu)"
    monkeypatch.setattr(t, "_fallback_to_cpu", fake_cpu)
    out = t(scout.Path("x.m4a"))
    assert out["backend"] == "whisper small (cpu)" and out["segments"] == [[0.0, 2.0, "hi there"]]

    class Broken:
        def transcribe(self, *a, **k):
            raise ValueError("bad file")
    t._m, t._kind = Broken(), "openai"
    with pytest.raises(ValueError):
        t(scout.Path("x.m4a"))


def test_cli_transcribe_flag():
    args = scout.build_parser().parse_args(["matcha", "--transcribe"])
    assert args.transcribe == "turbo"
    assert scout.build_parser().parse_args(["matcha", "--transcribe", "small"]).transcribe == "small"
    assert scout.build_parser().parse_args(["matcha"]).transcribe is None
